# -*- coding: utf-8 -*-
"""
===============================================================================
  HỆ THỐNG TẠO THÔNG BÁO TẠM NGỪNG CUNG CẤP ĐIỆN  (v2.0 - Multi-select)
  Dành cho ngành Điện lực Việt Nam - Tổng công ty Điện lực miền Bắc (EVNNPC)
-------------------------------------------------------------------------------
  Nâng cấp v2.0:
    - Chọn nhiều Lộ đường dây cùng lúc (multi-select)
    - Chọn nhiều Trạm biến áp cùng lúc (multi-select)
    - Có thể kết hợp cả Lộ + TBA trong cùng một thông báo
    - Dropdown tự sinh từ Excel - gõ để search nhanh
    - Hiển thị số liệu live khi chọn
-------------------------------------------------------------------------------
  Cách chạy:  streamlit run app.py
===============================================================================
"""

import io
import re
from datetime import datetime, time

import pandas as pd
import streamlit as st
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ====================================================================
# 1. CẤU HÌNH TRANG VÀ GIAO DIỆN THEO PHONG CÁCH EVN
# ====================================================================
st.set_page_config(
    page_title="Tạo Thông báo Cắt điện - EVNNPC",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded",
)

EVN_CSS = """
<style>
    :root {
        --evn-blue: #003D7A;
        --evn-light-blue: #0066CC;
        --evn-orange: #FF6B00;
        --evn-red:  #E30613;
    }
    html, body, [class*="css"] {
        font-family: 'Times New Roman', 'Segoe UI', sans-serif !important;
    }
    .main-header {
        background: linear-gradient(135deg, #003D7A 0%, #0066CC 100%);
        padding: 25px 20px; border-radius: 12px;
        color: white; text-align: center; margin-bottom: 25px;
        box-shadow: 0 4px 12px rgba(0,61,122,0.25);
    }
    .main-header h1 { margin: 0; font-size: 26px; font-weight: 800; }
    .main-header p { margin: 6px 0 0 0; font-size: 14px; opacity: 0.95; }
    .info-box {
        background-color: #E8F4FD; border-left: 4px solid #0066CC;
        padding: 12px 15px; margin: 10px 0; border-radius: 6px;
    }
    .warning-box {
        background-color: #FFF4E5; border-left: 4px solid #FF6B00;
        padding: 12px 15px; margin: 10px 0; border-radius: 6px;
    }
    .success-box {
        background-color: #E8F5E9; border-left: 4px solid #2E7D32;
        padding: 12px 15px; margin: 10px 0; border-radius: 6px;
    }
    .preview-box {
        background-color: #FFFFFF; border: 1px solid #DDD;
        padding: 30px 40px; border-radius: 8px;
        font-family: 'Times New Roman', serif; color: #000;
        max-height: 600px; overflow-y: auto;
        box-shadow: 0 2px 8px rgba(0,0,0,0.06);
    }
    .preview-box p { line-height: 1.6; margin: 8px 0; }
    .stButton>button {
        background: linear-gradient(135deg, #003D7A 0%, #0066CC 100%);
        color: white; font-weight: bold; border: none;
        padding: 12px 30px; border-radius: 6px; font-size: 16px;
    }
    .stButton>button:hover {
        background: linear-gradient(135deg, #0066CC 0%, #003D7A 100%);
        transform: translateY(-1px);
    }
    .stDownloadButton>button {
        background: linear-gradient(135deg, #FF6B00 0%, #E30613 100%) !important;
        color: white !important; font-weight: bold !important;
        padding: 12px 30px !important; font-size: 16px !important;
    }
    h2, h3, h4 { color: #003D7A; }
    [data-testid="stMetricValue"] { color: #003D7A; font-weight: bold; }
    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, #F0F6FC 0%, #E8F0F8 100%);
    }
    .stMultiSelect [data-baseweb="tag"] {
        background-color: #0066CC !important;
    }
</style>
"""
st.markdown(EVN_CSS, unsafe_allow_html=True)

st.markdown(
    """
<div class="main-header">
    <h1>⚡ HỆ THỐNG TẠO THÔNG BÁO TẠM NGỪNG CUNG CẤP ĐIỆN ⚡</h1>
    <p>TỔNG CÔNG TY ĐIỆN LỰC MIỀN BẮC – EVNNPC</p>
    <p style="font-size:12px;">Phiên bản 2.0 – Hỗ trợ chọn nhiều lộ / nhiều trạm cùng lúc</p>
</div>
""",
    unsafe_allow_html=True,
)


# ====================================================================
# 2. CÁC HÀM TIỆN ÍCH
# ====================================================================
def normalize_text(text) -> str:
    if pd.isna(text):
        return ""
    return str(text).strip().upper()


def find_column(df: pd.DataFrame, possible_names):
    """Tìm tên cột linh hoạt, không phân biệt hoa/thường."""
    cols_lower = {str(c).strip().lower(): c for c in df.columns}
    for name in possible_names:
        if name.strip().lower() in cols_lower:
            return cols_lower[name.strip().lower()]
    for name in possible_names:
        key = name.strip().lower()
        for c_low, c_orig in cols_lower.items():
            if key in c_low:
                return c_orig
    return None


@st.cache_data(show_spinner=False)
def load_excel(file_bytes: bytes) -> pd.DataFrame:
    """Đọc Excel từ bytes (cache lại để không phải đọc nhiều lần)."""
    df = pd.read_excel(io.BytesIO(file_bytes))
    df.columns = [str(c).strip() for c in df.columns]
    return df


def detect_columns(df: pd.DataFrame) -> dict:
    """Nhận diện các cột Lộ / TBA / Thôn / Xã / Huyện."""
    return {
        "lo":    find_column(df, ["Lộ", "Lộ đường dây", "Lo", "Lo duong day",
                                  "Feeder", "Đường dây", "Duong day"]),
        "tba":   find_column(df, ["TBA", "Trạm", "Trạm biến áp", "Tên TBA",
                                  "Ten TBA", "Tram bien ap", "Tram"]),
        "thon":  find_column(df, ["Thôn", "Xóm", "Thôn/Xóm", "Thon", "Xom"]),
        "xa":    find_column(df, ["Xã", "Phường", "Xã/Phường", "Xa", "Phuong"]),
        "huyen": find_column(df, ["Huyện", "Quận", "Huyen", "Quan"]),
    }


def filter_by_selection(df, selected_lo, selected_tba, col_map):
    """
    Lọc dữ liệu theo danh sách Lộ và/hoặc TBA được chọn.
    Logic: lấy tất cả dòng thỏa MỘT TRONG các điều kiện sau:
       - Cột Lộ nằm trong selected_lo
       - HOẶC cột TBA nằm trong selected_tba
    """
    col_lo = col_map.get("lo")
    col_tba = col_map.get("tba")
    col_thon = col_map.get("thon")
    col_xa = col_map.get("xa")
    col_huyen = col_map.get("huyen")

    if not selected_lo and not selected_tba:
        return df.iloc[0:0], [], [], []  # empty

    mask = pd.Series([False] * len(df), index=df.index)
    if selected_lo and col_lo:
        mask = mask | df[col_lo].astype(str).str.strip().isin(selected_lo)
    if selected_tba and col_tba:
        mask = mask | df[col_tba].astype(str).str.strip().isin(selected_tba)

    filtered = df[mask].copy()
    if filtered.empty:
        return filtered, [], [], []

    # Lộ và TBA thực tế xuất hiện trong dữ liệu đã lọc
    list_lo = sorted({str(x).strip() for x in filtered[col_lo].dropna() if str(x).strip()}) if col_lo else []
    list_tba = sorted({str(x).strip() for x in filtered[col_tba].dropna() if str(x).strip()}) if col_tba else []

    # Gom thôn/xã không trùng
    list_thon_xa = []
    if col_xa:
        grouped = {}
        cols_needed = [c for c in [col_thon, col_xa, col_huyen] if c]
        sub = filtered[cols_needed].dropna(subset=[col_xa])
        for _, row in sub.iterrows():
            xa = str(row[col_xa]).strip()
            huyen = str(row[col_huyen]).strip() if col_huyen else ""
            xa_key = f"{xa}|{huyen}" if huyen else xa
            if xa_key not in grouped:
                grouped[xa_key] = {"xa": xa, "huyen": huyen, "thons": set()}
            if col_thon:
                thon = str(row[col_thon]).strip()
                if thon and thon.lower() not in ("nan", "none", ""):
                    grouped[xa_key]["thons"].add(thon)

        for key in sorted(grouped.keys()):
            info = grouped[key]
            thons = sorted(info["thons"])
            huyen_text = f", huyện {info['huyen']}" if info["huyen"] else ""
            if thons:
                if len(thons) == 1:
                    list_thon_xa.append(f"thôn {thons[0]} – xã {info['xa']}{huyen_text}")
                else:
                    list_thon_xa.append(
                        f"các thôn {', '.join(thons)} – xã {info['xa']}{huyen_text}"
                    )
            else:
                list_thon_xa.append(f"xã {info['xa']}{huyen_text}")

    return filtered, list_lo, list_tba, list_thon_xa


# ====================================================================
# 3. XỬ LÝ FILE WORD
# ====================================================================
def _replace_in_paragraph(paragraph, replacements: dict):
    full_text = paragraph.text
    if not any(k in full_text for k in replacements):
        return
    new_text = full_text
    for k, v in replacements.items():
        new_text = new_text.replace(k, str(v))
    if new_text == full_text:
        return
    runs = paragraph.runs
    if not runs:
        return
    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""


def _replace_in_table(table, replacements: dict):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                _replace_in_paragraph(para, replacements)
            for inner_tbl in cell.tables:
                _replace_in_table(inner_tbl, replacements)


def replace_placeholders(doc: Document, replacements: dict):
    for para in doc.paragraphs:
        _replace_in_paragraph(para, replacements)
    for tbl in doc.tables:
        _replace_in_table(tbl, replacements)
    for section in doc.sections:
        for hdr in (section.header, section.first_page_header, section.even_page_header):
            if hdr is None:
                continue
            for para in hdr.paragraphs:
                _replace_in_paragraph(para, replacements)
            for tbl in hdr.tables:
                _replace_in_table(tbl, replacements)
        for ftr in (section.footer, section.first_page_footer, section.even_page_footer):
            if ftr is None:
                continue
            for para in ftr.paragraphs:
                _replace_in_paragraph(para, replacements)
            for tbl in ftr.tables:
                _replace_in_table(tbl, replacements)


def ensure_vietnamese_font(doc: Document, font_name="Times New Roman"):
    def _set_run_font(run):
        run.font.name = font_name
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), font_name)
        rFonts.set(qn("w:hAnsi"), font_name)
        rFonts.set(qn("w:cs"), font_name)
        rFonts.set(qn("w:eastAsia"), font_name)

    def _process_paragraphs(paras):
        for p in paras:
            for r in p.runs:
                _set_run_font(r)

    _process_paragraphs(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                _process_paragraphs(cell.paragraphs)


def format_date_vn(d) -> str:
    return f"ngày {d.day:02d} tháng {d.month:02d} năm {d.year}"


def format_time_vn(t) -> str:
    return f"{t.hour:02d}h{t.minute:02d}"


# ====================================================================
# 4. SIDEBAR - UPLOAD FILES
# ====================================================================
with st.sidebar:
    st.markdown("### 📁 TẢI FILE LÊN")
    excel_file = st.file_uploader(
        "📊 File Excel dữ liệu khách hàng",
        type=["xlsx", "xls"],
        help="File Excel có các cột: Lộ đường dây, TBA, Thôn, Xã, Huyện",
    )
    word_file = st.file_uploader(
        "📄 File Word mẫu thông báo (.docx)",
        type=["docx"],
        help="File Word chứa các placeholder dạng {{NGAY_CAT}}, {{LO_DUONG_DAY}}, ...",
    )

    st.markdown("---")
    st.markdown("### 🔑 PLACEHOLDER WORD")
    st.code(
        """{{NGAY_CAT}}
{{GIO_BAT_DAU}}
{{GIO_KET_THUC}}
{{LO_DUONG_DAY}}
{{DS_TRAM}}
{{DS_THON_XA}}
{{LY_DO}}""",
        language="text",
    )

    st.markdown("---")
    with st.expander("📖 Cấu trúc file Excel"):
        st.markdown(
            """
**Cột Excel cần có (tên linh hoạt):**

| Cột | Tên gợi ý |
|-----|-----------|
| **Lộ đường dây** ✅ | `Lộ`, `Feeder`, `Đường dây` |
| **TBA** ✅ | `TBA`, `Trạm biến áp` |
| **Thôn** | `Thôn`, `Xóm` |
| **Xã** ✅ | `Xã`, `Phường` |
| **Huyện** | `Huyện` *(tùy chọn)* |
            """
        )


# ====================================================================
# 5. KHU VỰC CHÍNH - PHỤ THUỘC VÀO FILE ĐÃ UPLOAD
# ====================================================================
if not excel_file:
    st.markdown(
        '<div class="warning-box">⬅️ <b>Vui lòng tải lên file Excel dữ liệu khách hàng</b> '
        'ở thanh bên trái để bắt đầu.</div>',
        unsafe_allow_html=True,
    )
    st.stop()

# Đọc Excel và phát hiện cột
try:
    excel_bytes = excel_file.getvalue()
    df = load_excel(excel_bytes)
    col_map = detect_columns(df)
except Exception as e:
    st.error(f"❌ Lỗi đọc file Excel: {e}")
    st.stop()

# Kiểm tra cột bắt buộc
missing = []
if not col_map["lo"]:
    missing.append("Lộ đường dây")
if not col_map["tba"]:
    missing.append("Trạm biến áp")
if not col_map["xa"]:
    missing.append("Xã/Phường")
if missing:
    st.error(f"❌ File Excel thiếu các cột bắt buộc: **{', '.join(missing)}**")
    st.info(f"Các cột hiện có trong Excel: `{', '.join(df.columns)}`")
    st.stop()

# Lấy danh sách duy nhất cho dropdown
all_lo = sorted({str(x).strip() for x in df[col_map["lo"]].dropna() if str(x).strip()})
all_tba = sorted({str(x).strip() for x in df[col_map["tba"]].dropna() if str(x).strip()})

# Thông tin tóm tắt Excel
c1, c2, c3, c4 = st.columns(4)
c1.metric("📋 Tổng bản ghi", f"{len(df):,}")
c2.metric("🔌 Số Lộ ĐD", len(all_lo))
c3.metric("🏭 Số TBA", len(all_tba))
c4.metric("📄 Word mẫu", "✅ OK" if word_file else "❌ Thiếu")

st.markdown("---")

# ====================================================================
# 6. PHẦN CHỌN LỘ / TBA - MULTI-SELECT
# ====================================================================
st.markdown("### 🎯 CHỌN KHU VỰC BỊ ẢNH HƯỞNG")
st.markdown(
    '<div class="info-box">💡 <b>Có thể chọn nhiều Lộ và nhiều TBA cùng lúc.</b> '
    'Gõ phím trong ô để search nhanh. Kết quả sẽ là hợp của tất cả các lựa chọn.</div>',
    unsafe_allow_html=True,
)

# Khởi tạo session_state nếu chưa có
if "multi_lo" not in st.session_state:
    st.session_state["multi_lo"] = []
if "multi_tba" not in st.session_state:
    st.session_state["multi_tba"] = []

# Callback functions: chạy TRƯỚC khi widget được render -> tránh lỗi
def _select_all_lo():
    st.session_state["multi_lo"] = list(all_lo)

def _clear_all_lo():
    st.session_state["multi_lo"] = []

def _select_all_tba():
    st.session_state["multi_tba"] = list(all_tba)

def _clear_all_tba():
    st.session_state["multi_tba"] = []

col_lo_select, col_tba_select = st.columns(2)

with col_lo_select:
    st.markdown("#### 🔌 Chọn Lộ đường dây")

    # Đặt nút thao tác nhanh LÊN TRÊN multiselect, dùng on_click callback
    cc1, cc2 = st.columns(2)
    with cc1:
        st.button(
            "✅ Chọn tất cả lộ",
            use_container_width=True,
            key="btn_all_lo",
            on_click=_select_all_lo,
        )
    with cc2:
        st.button(
            "❌ Bỏ chọn tất cả lộ",
            use_container_width=True,
            key="btn_clear_lo",
            on_click=_clear_all_lo,
        )

    selected_lo = st.multiselect(
        "Chọn một hoặc nhiều lộ:",
        options=all_lo,
        placeholder="Bấm để chọn hoặc gõ để search...",
        label_visibility="collapsed",
        key="multi_lo",
    )
    if selected_lo:
        st.caption(f"✅ Đã chọn **{len(selected_lo)}** / {len(all_lo)} lộ")
    else:
        st.caption(f"Chưa chọn lộ nào (tổng có **{len(all_lo)}** lộ)")

with col_tba_select:
    st.markdown("#### 🏭 Chọn Trạm biến áp")

    cc1, cc2 = st.columns(2)
    with cc1:
        st.button(
            "✅ Chọn tất cả TBA",
            use_container_width=True,
            key="btn_all_tba",
            on_click=_select_all_tba,
        )
    with cc2:
        st.button(
            "❌ Bỏ chọn tất cả TBA",
            use_container_width=True,
            key="btn_clear_tba",
            on_click=_clear_all_tba,
        )

    selected_tba = st.multiselect(
        "Chọn một hoặc nhiều TBA:",
        options=all_tba,
        placeholder="Bấm để chọn hoặc gõ để search...",
        label_visibility="collapsed",
        key="multi_tba",
    )
    if selected_tba:
        st.caption(f"✅ Đã chọn **{len(selected_tba)}** / {len(all_tba)} TBA")
    else:
        st.caption(f"Chưa chọn TBA nào (tổng có **{len(all_tba)}** TBA)")

# ====================================================================
# 7. PREVIEW SỐ LIỆU LIVE (chưa cần bấm submit)
# ====================================================================
if selected_lo or selected_tba:
    with st.spinner("🔄 Đang lọc dữ liệu..."):
        filtered_preview, list_lo_p, list_tba_p, list_thon_xa_p = filter_by_selection(
            df, selected_lo, selected_tba, col_map
        )

    st.markdown("#### 📊 Số liệu khi áp dụng các lựa chọn:")
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("🔌 Tổng lộ liên quan", len(list_lo_p))
    m2.metric("🏭 Tổng TBA liên quan", len(list_tba_p))
    m3.metric("🏘️ Khu vực (xã)", len(list_thon_xa_p))
    m4.metric("📋 Bản ghi", len(filtered_preview))

    with st.expander("👁️ Xem danh sách chi tiết sẽ xuất hiện trong thông báo"):
        cA, cB = st.columns(2)
        with cA:
            st.markdown(f"**🔌 Lộ đường dây ({len(list_lo_p)}):**")
            for x in list_lo_p:
                st.markdown(f"- {x}")
            st.markdown(f"**🏭 Trạm biến áp ({len(list_tba_p)}):**")
            for x in list_tba_p:
                st.markdown(f"- {x}")
        with cB:
            st.markdown(f"**🏘️ Khu vực ảnh hưởng ({len(list_thon_xa_p)}):**")
            for x in list_thon_xa_p:
                st.markdown(f"- {x}")
else:
    st.markdown(
        '<div class="warning-box">⚠️ Hãy chọn ít nhất một Lộ hoặc một TBA ở trên để tiếp tục.</div>',
        unsafe_allow_html=True,
    )

st.markdown("---")

# ====================================================================
# 8. FORM NHẬP THÔNG TIN THÔNG BÁO
# ====================================================================
st.markdown("### 📝 THÔNG TIN THÔNG BÁO")

with st.form("form_thong_bao", clear_on_submit=False):
    c1, c2, c3 = st.columns(3)
    with c1:
        ngay_cat = st.date_input("📅 Ngày cắt điện", value=datetime.now())
    with c2:
        gio_bat_dau = st.time_input("🕐 Giờ bắt đầu", value=time(7, 0))
    with c3:
        gio_ket_thuc = st.time_input("🕔 Giờ kết thúc", value=time(17, 0))

    ly_do = st.text_area(
        "📋 Lý do công tác",
        value="Sửa chữa, bảo dưỡng định kỳ lưới điện nhằm đảm bảo an toàn và "
              "nâng cao chất lượng cung cấp điện",
        height=80,
    )

    submitted = st.form_submit_button(
        "⚡ TẠO THÔNG BÁO ⚡", use_container_width=True
    )


# ====================================================================
# 9. XỬ LÝ KHI BẤM NÚT [TẠO THÔNG BÁO]
# ====================================================================
if submitted:
    if not word_file:
        st.error("❌ Vui lòng tải lên file **Word mẫu thông báo** ở thanh bên trái!")
        st.stop()
    if not selected_lo and not selected_tba:
        st.error("❌ Vui lòng chọn ít nhất một **Lộ** hoặc một **TBA**!")
        st.stop()
    if gio_ket_thuc <= gio_bat_dau:
        st.warning("⚠️ Giờ kết thúc nên muộn hơn giờ bắt đầu. Vui lòng kiểm tra lại.")

    try:
        # Lọc dữ liệu
        filtered, list_lo, list_tba, list_thon_xa = filter_by_selection(
            df, selected_lo, selected_tba, col_map
        )
        if filtered.empty:
            st.warning("⚠️ Không tìm thấy dữ liệu nào với lựa chọn này.")
            st.stop()

        # Chuẩn bị placeholder
        ds_tram_str = "; ".join(list_tba)
        ds_thon_xa_str = "; ".join(list_thon_xa)
        ds_lo_str = ", ".join(list_lo)

        replacements = {
            "{{NGAY_CAT}}":     format_date_vn(ngay_cat),
            "{{GIO_BAT_DAU}}":  format_time_vn(gio_bat_dau),
            "{{GIO_KET_THUC}}": format_time_vn(gio_ket_thuc),
            "{{LO_DUONG_DAY}}": ds_lo_str,
            "{{DS_TRAM}}":      ds_tram_str,
            "{{DS_THON_XA}}":   ds_thon_xa_str,
            "{{LY_DO}}":        ly_do.strip(),
        }

        # Preview nội dung
        st.markdown("### 👁️ XEM TRƯỚC NỘI DUNG THÔNG BÁO")
        preview_html = f"""
        <div class="preview-box">
            <table style="width:100%; border:none;">
                <tr>
                    <td style="width:50%; text-align:center; vertical-align:top; border:none;">
                        <b>TỔNG CÔNG TY ĐIỆN LỰC MIỀN BẮC</b><br>
                        <b style="text-decoration:underline;">CÔNG TY ĐIỆN LỰC ...</b>
                    </td>
                    <td style="width:50%; text-align:center; vertical-align:top; border:none;">
                        <b>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br>
                        <b style="text-decoration:underline;">Độc lập – Tự do – Hạnh phúc</b>
                    </td>
                </tr>
            </table>
            <p style="text-align:right; font-style:italic;">............, {format_date_vn(ngay_cat)}</p>

            <h3 style="text-align:center; margin-top:20px;">THÔNG BÁO</h3>
            <p style="text-align:center;"><b><i>V/v tạm ngừng cung cấp điện</i></b></p>

            <p><b>Kính gửi:</b> Quý khách hàng sử dụng điện.</p>

            <p style="text-indent: 30px;">Thực hiện kế hoạch công tác trên lưới điện nhằm đảm bảo
            vận hành an toàn, ổn định và nâng cao chất lượng cung cấp điện phục vụ Quý khách hàng,
            Công ty Điện lực trân trọng thông báo lịch tạm ngừng cung cấp điện như sau:</p>

            <p><b>1. Thời gian tạm ngừng cấp điện:</b><br>
            &nbsp;&nbsp;&nbsp;&nbsp;Từ <b>{format_time_vn(gio_bat_dau)}</b>
            đến <b>{format_time_vn(gio_ket_thuc)}</b> {format_date_vn(ngay_cat)}.</p>

            <p><b>2. Lộ đường dây liên quan ({len(list_lo)} lộ):</b><br>
            &nbsp;&nbsp;&nbsp;&nbsp;{ds_lo_str}</p>

            <p><b>3. Trạm biến áp bị ảnh hưởng ({len(list_tba)} TBA):</b><br>
            &nbsp;&nbsp;&nbsp;&nbsp;{ds_tram_str}</p>

            <p><b>4. Khu vực mất điện ({len(list_thon_xa)} khu vực):</b><br>
            &nbsp;&nbsp;&nbsp;&nbsp;{ds_thon_xa_str}.</p>

            <p><b>5. Lý do:</b> {ly_do}.</p>

            <p style="text-indent: 30px;">Công ty Điện lực rất mong Quý khách hàng thông cảm
            và chủ động bố trí công việc, sinh hoạt, sản xuất phù hợp trong thời gian tạm
            ngừng cấp điện nêu trên.</p>

            <p style="text-indent: 30px;">Mọi thông tin chi tiết, Quý khách hàng vui lòng
            liên hệ Tổng đài CSKH <b>19006769</b> hoặc website <b>cskh.npc.com.vn</b>.</p>

            <p style="text-indent: 30px;"><i>Trân trọng thông báo!</i></p>
        </div>
        """
        st.markdown(preview_html, unsafe_allow_html=True)

        st.markdown(
            '<div class="info-box">📌 <b>Lưu ý:</b> Phần preview trên chỉ minh họa nội dung. '
            'File Word xuất ra sẽ <b>giữ nguyên định dạng gốc</b> của file mẫu bạn upload.</div>',
            unsafe_allow_html=True,
        )

        # Tạo Word
        with st.spinner("📄 Đang tạo file Word..."):
            word_file.seek(0)
            doc = Document(word_file)
            replace_placeholders(doc, replacements)
            ensure_vietnamese_font(doc, "Times New Roman")

            output = io.BytesIO()
            doc.save(output)
            output.seek(0)

        # Tên file thông minh
        if len(selected_lo) == 1 and not selected_tba:
            tag = re.sub(r"[^\w\-]", "_", selected_lo[0])
        elif len(selected_tba) == 1 and not selected_lo:
            tag = re.sub(r"[^\w\-]", "_", selected_tba[0])
        elif selected_lo and not selected_tba:
            tag = f"{len(selected_lo)}Lo"
        elif selected_tba and not selected_lo:
            tag = f"{len(selected_tba)}TBA"
        else:
            tag = f"{len(selected_lo)}Lo_{len(selected_tba)}TBA"

        file_name = f"ThongBao_CatDien_{ngay_cat.strftime('%d-%m-%Y')}_{tag}.docx"

        st.markdown("### 📥 TẢI FILE WORD")
        st.markdown(
            '<div class="success-box">✅ <b>Tạo thông báo thành công!</b> '
            'Bấm nút bên dưới để tải file Word về máy.</div>',
            unsafe_allow_html=True,
        )

        st.download_button(
            label="⬇️ TẢI FILE WORD THÔNG BÁO",
            data=output,
            file_name=file_name,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

        with st.expander("📊 Xem bảng dữ liệu Excel đã lọc"):
            st.dataframe(filtered, use_container_width=True, height=300)

    except Exception as e:
        st.error(f"❌ Có lỗi xảy ra: **{str(e)}**")
        import traceback
        with st.expander("🔧 Chi tiết kỹ thuật"):
            st.code(traceback.format_exc())


# ====================================================================
# 10. FOOTER
# ====================================================================
st.markdown("---")
st.markdown(
    """
<div style="text-align:center; color:#555; font-size:13px; padding:15px;">
    <b>⚡ HỆ THỐNG TỰ ĐỘNG HÓA THÔNG BÁO CẮT ĐIỆN – v2.0 ⚡</b><br>
    Phát triển dành cho ngành Điện lực Việt Nam – EVNNPC<br>
    <i>Powered by Python · Streamlit · python-docx · pandas</i>
</div>
""",
    unsafe_allow_html=True,
)
